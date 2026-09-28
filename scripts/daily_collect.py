#!/usr/bin/env python3
"""
Daily news collection pipeline.

Orchestrates the full daily collection workflow:
1. Fetch from all sources (fetch_news.py)
2. AI filter + summarize (ai_filter.py)
3. Merge multi-source articles (ai_filter.py)
4. Dedup against history (dedup.py)
5. Download images (fetch_image.py)
6. Generate candidates.json, candidates.xlsx, candidates.html
7. Push to cloud (push_to_cloud.py)
8. Update history (dedup.py)

Usage:
    python scripts/daily_collect.py [--days-back 3] [--skip-push]
"""

import json
import logging
import os
import sys
from datetime import datetime
from pathlib import Path

# Ensure scripts dir is on path for imports
SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SCRIPT_DIR.parent
sys.path.insert(0, str(SCRIPT_DIR))

from fetch_news import fetch_all_sources, SOURCES_PATH
from ai_filter import filter_and_summarize, merge_multi_source
from dedup import dedup_candidates, update_history
from fetch_image import fetch_article_image
from push_to_cloud import push_candidates, reset_confirm_queue, push_to_gist

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
)
logger = logging.getLogger(__name__)


def _generate_xlsx(candidates: list[dict], output_path: Path):
    """Generate an Excel file from candidates for easy review."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Candidates"

        # Header style
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        headers = ["ID", "中文标题", "英文标题", "来源", "发布日期", "摘要", "关键实体", "链接"]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        # Data rows
        for i, c in enumerate(candidates, 1):
            ws.cell(row=i + 1, column=1, value=i)
            ws.cell(row=i + 1, column=2, value=c.get("title_cn", ""))
            ws.cell(row=i + 1, column=3, value=c.get("title_en", ""))
            ws.cell(row=i + 1, column=4, value=c.get("source", ""))
            ws.cell(row=i + 1, column=5, value=c.get("pub_date", ""))
            ws.cell(row=i + 1, column=6, value=c.get("summary_cn", ""))
            ws.cell(row=i + 1, column=7, value=", ".join(c.get("key_entities", [])))
            urls = c.get("source_urls", [])
            ws.cell(row=i + 1, column=8, value=urls[0] if urls else "")

        # Column widths
        ws.column_dimensions["A"].width = 5
        ws.column_dimensions["B"].width = 35
        ws.column_dimensions["C"].width = 40
        ws.column_dimensions["D"].width = 20
        ws.column_dimensions["E"].width = 12
        ws.column_dimensions["F"].width = 50
        ws.column_dimensions["G"].width = 30
        ws.column_dimensions["H"].width = 50

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output_path))
        logger.info("Generated xlsx: %s", output_path)

    except ImportError:
        logger.warning("openpyxl not installed, skipping xlsx generation")
    except Exception as e:
        logger.error("Failed to generate xlsx: %s", e)


def _generate_html(candidates: list[dict], output_path: Path, date_str: str):
    """Generate an HTML page for candidate review."""
    rows = []
    for i, c in enumerate(candidates, 1):
        urls = c.get("source_urls", [])
        url_html = ""
        for u in urls:
            url_html += f'<a href="{u}" target="_blank">{u}</a><br>'
        entities = ", ".join(c.get("key_entities", []))
        rows.append(f"""
        <tr>
            <td>{i}</td>
            <td><strong>{c.get('title_cn', '')}</strong></td>
            <td>{c.get('title_en', '')}</td>
            <td>{c.get('source', '')}</td>
            <td>{c.get('pub_date', '')}</td>
            <td>{c.get('summary_cn', '')}</td>
            <td>{entities}</td>
            <td>{url_html}</td>
        </tr>""")

    html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>橡胶行业新闻候选 - {date_str}</title>
    <style>
        body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; margin: 20px; background: #f5f5f5; }}
        h1 {{ color: #333; }}
        table {{ border-collapse: collapse; width: 100%; background: white; box-shadow: 0 1px 3px rgba(0,0,0,0.12); }}
        th {{ background: #4472C4; color: white; padding: 10px 8px; text-align: left; font-size: 13px; }}
        td {{ padding: 8px; border-bottom: 1px solid #eee; font-size: 13px; vertical-align: top; }}
        tr:hover {{ background: #f0f4ff; }}
        a {{ color: #4472C4; word-break: break-all; font-size: 12px; }}
        .meta {{ color: #666; font-size: 12px; margin-bottom: 15px; }}
    </style>
</head>
<body>
    <h1>橡胶行业新闻候选 - {date_str}</h1>
    <p class="meta">共 {len(candidates)} 条候选新闻 | 生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M')}</p>
    <table>
        <thead>
            <tr>
                <th>#</th>
                <th>中文标题</th>
                <th>英文标题</th>
                <th>来源</th>
                <th>日期</th>
                <th>摘要</th>
                <th>关键实体</th>
                <th>链接</th>
            </tr>
        </thead>
        <tbody>
            {''.join(rows)}
        </tbody>
    </table>
</body>
</html>"""

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(html, encoding="utf-8")
    logger.info("Generated HTML: %s", output_path)


def main():
    import argparse
    parser = argparse.ArgumentParser(description="Daily news collection pipeline")
    parser.add_argument("--days-back", type=int, default=3, help="Look back N days for sources")
    parser.add_argument("--skip-push", action="store_true", help="Skip pushing to cloud storage")
    parser.add_argument("--skip-images", action="store_true", help="Skip image downloading")
    args = parser.parse_args()

    date_str = datetime.now().strftime("%Y-%m-%d")
    data_dir = PROJECT_DIR / "data" / date_str
    data_dir.mkdir(parents=True, exist_ok=True)

    logger.info("=== Daily Collection Pipeline: %s ===", date_str)

    # Step 1: Fetch from all sources
    logger.info("Step 1: Fetching from all sources...")
    # In GitHub Actions, exclude ERJ (Cloudflare blocks datacenter IPs).
    # ERJ is handled by local script when computer is on.
    exclude_sources = os.environ.get("EXCLUDE_SOURCES", "").split(",") if os.environ.get("EXCLUDE_SOURCES") else []
    raw_articles = fetch_all_sources(SOURCES_PATH, days_back=args.days_back, exclude=exclude_sources)
    if not raw_articles:
        logger.warning("No articles fetched from any source. Exiting.")
        return

    raw_path = data_dir / "raw_articles.json"
    with open(raw_path, "w", encoding="utf-8") as f:
        json.dump(raw_articles, f, ensure_ascii=False, indent=2)
    logger.info("Fetched %d raw articles", len(raw_articles))

    # Step 2: AI filter + summarize
    logger.info("Step 2: AI filtering and summarizing...")
    candidates = filter_and_summarize(raw_articles)
    logger.info("After filtering: %d candidates", len(candidates))

    # Step 3: Merge multi-source articles
    logger.info("Step 3: Merging multi-source duplicates...")
    candidates = merge_multi_source(candidates)
    logger.info("After merging: %d unique stories", len(candidates))

    # Step 4: Dedup against history
    logger.info("Step 4: Deduplicating against history...")
    history_path = PROJECT_DIR / "data" / "selected_history.json"
    candidates = dedup_candidates(candidates, history_path)
    logger.info("After dedup: %d new candidates", len(candidates))

    if not candidates:
        logger.info("No new candidates after dedup. Nothing to do.")
        return

    # Assign IDs
    for i, c in enumerate(candidates, 1):
        c["id"] = i

    # Step 5: Download images
    if not args.skip_images:
        logger.info("Step 5: Downloading images...")
        for c in candidates:
            urls = c.get("source_urls", [])
            if urls:
                try:
                    image_path = fetch_article_image(urls[0], c.get("source_id", "unknown"), date_str)
                    if image_path:
                        c["image_url"] = image_path
                except Exception as e:
                    logger.warning("Image fetch failed for '%s': %s", c.get("title_cn", "?")[:30], e)
    else:
        logger.info("Step 5: Skipping image download")

    # Step 6: Generate output files
    logger.info("Step 6: Generating output files...")

    # candidates.json
    candidates_path = data_dir / "candidates.json"
    with open(candidates_path, "w", encoding="utf-8") as f:
        json.dump(candidates, f, ensure_ascii=False, indent=2)

    # candidates.xlsx
    xlsx_path = data_dir / "candidates.xlsx"
    _generate_xlsx(candidates, xlsx_path)

    # candidates.html
    html_path = data_dir / "candidates.html"
    _generate_html(candidates, html_path, date_str)

    logger.info("Generated %d candidates in %s", len(candidates), data_dir)

    # Step 7: Push to cloud
    if not args.skip_push:
        logger.info("Step 7: Pushing to cloud storage...")
        push_candidates(candidates, date_str)
        reset_confirm_queue(candidates, date_str)
        push_to_gist(candidates)
    else:
        logger.info("Step 7: Skipping cloud push")

    # Step 8: Update history
    logger.info("Step 8: Updating history...")
    update_history(candidates, history_path, date_str)

    logger.info("=== Daily Collection Complete: %d candidates ===", len(candidates))


if __name__ == "__main__":
    main()
